"""Google-Sheets W2C spreadsheet download and parsing."""

import io
import json
import logging
from dataclasses import dataclass

import httpx
import openpyxl

from .config import MODEL, USER_AGENT
from .extract import resolve_product_link

logger = logging.getLogger(__name__)

MAX_TABS = 20
MAX_ROWS_PER_TAB = 2000
EXPORT_URL = "https://docs.google.com/spreadsheets/d/{key}/export?format=xlsx"

COLUMN_MAP_PROMPT = """You are looking at the first rows of one tab of a community fashion-item spreadsheet. Decide whether this tab lists purchasable items, and if so which 0-based column index holds each field.

Respond with ONLY a JSON object, no other text:
{{"is_items": boolean, "name_col": int or null, "price_col": int or null, "link_col": int or null, "image_col": int or null, "currency": "CNY" | "USD" or null}}

`link_col` is the column with purchase links (weidian/taobao/agent links). A tab without both an item-name column and a link column is not an items tab.

TAB NAME: {tab_name}

FIRST ROWS (each line is one row, cells joined by " | "):
{preview}
"""


@dataclass
class ColumnMap:
    name_col: int
    link_col: int
    price_col: int | None
    image_col: int | None
    currency: str | None


def download_workbook(sheet_key: str, client: httpx.Client) -> openpyxl.Workbook:
    resp = client.get(
        EXPORT_URL.format(key=sheet_key),
        headers={"User-Agent": USER_AGENT},
        follow_redirects=True,
        timeout=60,
    )
    resp.raise_for_status()
    # read_only=False: read-only worksheets don't expose cell.hyperlink,
    # which parse_tab needs for display-text link cells.
    return openpyxl.load_workbook(io.BytesIO(resp.content), read_only=False, data_only=True)


def _extract_text(message) -> str:
    content = getattr(message, "content", None)
    if not content:
        raise ValueError("empty or non-text response")
    text = getattr(content[0], "text", None)
    if not isinstance(text, str):
        raise ValueError("empty or non-text response")
    return text


def parse_column_map(text: str) -> ColumnMap | None:
    start, end = text.find("{"), text.rfind("}")
    if start == -1 or end <= start:
        raise ValueError("no JSON object in column-map response")
    data = json.loads(text[start : end + 1])
    if not isinstance(data.get("is_items"), bool):
        raise ValueError("is_items missing")
    if not data["is_items"]:
        return None
    name_col, link_col = data.get("name_col"), data.get("link_col")
    if not isinstance(name_col, int) or not isinstance(link_col, int):
        return None
    currency = data.get("currency")
    return ColumnMap(
        name_col=name_col,
        link_col=link_col,
        price_col=data.get("price_col") if isinstance(data.get("price_col"), int) else None,
        image_col=data.get("image_col") if isinstance(data.get("image_col"), int) else None,
        currency=currency if currency in ("CNY", "USD") else None,
    )


def map_columns(llm, tab_name: str, preview_rows: list[list]) -> ColumnMap | None:
    preview = "\n".join(
        " | ".join("" if c is None else str(c) for c in row) for row in preview_rows[:10]
    )
    prompt = COLUMN_MAP_PROMPT.format(tab_name=tab_name, preview=preview[:4000])
    last_err: Exception | None = None
    for _ in range(2):  # initial attempt + 1 retry
        message = llm.messages.create(
            model=MODEL, max_tokens=256, messages=[{"role": "user", "content": prompt}]
        )
        try:
            return parse_column_map(_extract_text(message))
        except (ValueError, json.JSONDecodeError) as e:
            last_err = e
    logger.warning("column mapping unparseable for tab %s: %s", tab_name, last_err)
    return None


def _cell_link(ws_row, col: int) -> str:
    cell = ws_row[col] if col < len(ws_row) else None
    if cell is None:
        return ""
    link = getattr(cell, "hyperlink", None)
    if link is not None and getattr(link, "target", None):
        return str(link.target)
    return "" if cell.value is None else str(cell.value)


def _cell_text(ws_row, col: int | None) -> str | None:
    if col is None or col >= len(ws_row):
        return None
    value = ws_row[col].value
    return None if value is None else str(value).strip() or None


def parse_tab(ws, colmap: ColumnMap) -> list[dict]:
    rows: list[dict] = []
    for row_number, ws_row in enumerate(ws.iter_rows(), start=1):
        if row_number == 1:
            continue  # header
        if len(rows) >= MAX_ROWS_PER_TAB:
            break
        name = _cell_text(ws_row, colmap.name_col)
        raw_link = _cell_link(ws_row, colmap.link_col).strip()
        if not name and not raw_link:
            continue
        resolved = resolve_product_link(raw_link)
        rows.append({
            "row_number": row_number,
            "name": name,
            "price_raw": _cell_text(ws_row, colmap.price_col),
            "currency": colmap.currency,
            "image_url": _cell_text(ws_row, colmap.image_col),
            "raw_link": raw_link or None,
            "product_url": resolved[1] if resolved else None,
            "platform": resolved[0] if resolved else None,
        })
    return rows
