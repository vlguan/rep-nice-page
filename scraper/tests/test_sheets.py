import io
import types

import openpyxl
import pytest

from scraper.sheets import ColumnMap, MAX_ROWS_PER_TAB, map_columns, parse_tab
from tests.conftest import requires_db


def make_ws(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    return ws


def test_parse_tab_extracts_rows_and_resolves_links():
    ws = make_ws([
        ["Item", "Price", "Link", "Pic"],
        ["Hellstar hoodie", "¥199", "https://weidian.com/item.html?itemID=11", "https://img/1.jpg"],
        ["AMIRI jeans", "25.5", "https://cnfans.com/product/?shop_type=taobao&id=22", "https://img/2.jpg"],
        ["No link row", "10", "", ""],
    ])
    cm = ColumnMap(name_col=0, price_col=1, link_col=2, image_col=3, currency="USD")
    rows = parse_tab(ws, cm)
    assert rows[0] == {
        "row_number": 2, "name": "Hellstar hoodie", "price_raw": "¥199", "currency": "USD",
        "image_url": "https://img/1.jpg", "raw_link": "https://weidian.com/item.html?itemID=11",
        "product_url": "https://weidian.com/item.html?itemID=11", "platform": "weidian",
    }
    assert rows[1]["platform"] == "taobao"
    assert rows[1]["product_url"] == "https://item.taobao.com/item.htm?id=22"
    assert rows[2]["product_url"] is None
    assert len(rows) == 3  # header skipped, empty name+link rows dropped


def test_parse_tab_uses_hyperlink_target_when_cell_is_display_text():
    ws = make_ws([["Item", "Link"], ["hoodie", "click here"]])
    ws.cell(row=2, column=2).hyperlink = "https://weidian.com/item.html?itemID=99"
    cm = ColumnMap(name_col=0, price_col=None, link_col=1, image_col=None, currency=None)
    rows = parse_tab(ws, cm)
    assert rows[0]["product_url"] == "https://weidian.com/item.html?itemID=99"


def test_parse_tab_caps_rows():
    data = [["Item", "Link"]] + [[f"i{n}", "https://weidian.com/item.html?itemID=1"] for n in range(MAX_ROWS_PER_TAB + 50)]
    ws = make_ws(data)
    cm = ColumnMap(name_col=0, price_col=None, link_col=1, image_col=None, currency=None)
    assert len(parse_tab(ws, cm)) == MAX_ROWS_PER_TAB


class FakeLLM:
    def __init__(self, text):
        self._text = text
        self.messages = self

    def create(self, **kwargs):
        return types.SimpleNamespace(content=[types.SimpleNamespace(text=self._text)])


def test_map_columns_parses_response():
    llm = FakeLLM('{"is_items": true, "name_col": 0, "price_col": 1, "link_col": 2, "image_col": null, "currency": "CNY"}')
    cm = map_columns(llm, "Hoodies", [["Item", "Price", "Link"]])
    assert cm == ColumnMap(name_col=0, price_col=1, link_col=2, image_col=None, currency="CNY")


def test_map_columns_not_items_tab_returns_none():
    llm = FakeLLM('{"is_items": false, "name_col": null, "price_col": null, "link_col": null, "image_col": null, "currency": null}')
    assert map_columns(llm, "Intro", [["Welcome to the sheet"]]) is None


def test_map_columns_requires_name_and_link():
    llm = FakeLLM('{"is_items": true, "name_col": 0, "price_col": null, "link_col": null, "image_col": null, "currency": null}')
    assert map_columns(llm, "Tab", [["Item"]]) is None


@requires_db
def test_sync_spreadsheets_stages_rows(conn):
    import httpx

    from scraper import db
    from scraper.sheets import sync_spreadsheets

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoodies"
    ws.append(["Item", "Price", "Link"])
    ws.append(["Hellstar hoodie", "199", "https://weidian.com/item.html?itemID=77"])
    buf = io.BytesIO()
    wb.save(buf)

    def handler(request):
        return httpx.Response(200, content=buf.getvalue())

    client = httpx.Client(transport=httpx.MockTransport(handler))
    llm = FakeLLM('{"is_items": true, "name_col": 0, "price_col": 1, "link_col": 2, "image_col": null, "currency": "CNY"}')

    db.upsert_spreadsheet(conn, "sync-key", "https://docs.google.com/spreadsheets/d/sync-key", None)
    sync_spreadsheets(conn, llm, client)

    row = conn.execute("SELECT name, product_url, platform FROM spreadsheet_rows").fetchone()
    assert row == ("Hellstar hoodie", "https://weidian.com/item.html?itemID=77", "weidian")
    assert conn.execute("SELECT last_synced_at FROM spreadsheets").fetchone()[0] is not None


@requires_db
def test_sync_spreadsheets_records_error(conn):
    import httpx

    from scraper import db
    from scraper.sheets import sync_spreadsheets

    def handler(request):
        return httpx.Response(403)

    client = httpx.Client(transport=httpx.MockTransport(handler))
    db.upsert_spreadsheet(conn, "denied-key", "u", None)
    sync_spreadsheets(conn, None, client)  # llm never reached
    err = conn.execute("SELECT sync_error FROM spreadsheets").fetchone()[0]
    assert err and "403" in err
